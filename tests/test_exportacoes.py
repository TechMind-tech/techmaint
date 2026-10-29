"""
Exportacoes PDF / XLSX / CSV (SPEC §4: /dashboard/exportar.pdf|xlsx e /relatorios/<aba>/exportar.pdf|xlsx|csv).

Rastreabilidade: RF "Emissao de Relatorios ... exportacao em PDF, Excel e CSV" (Projeto §15b), Projeto §8 ("Exportacao de relatorios
em Excel/PDF"), Manual 3.1/3.2 (botoes Exportar PDF/Excel), SPEC §4 (cabecalho TechMaint, filtros aplicados e data).
"""

import csv
import datetime
import io
import re
import shutil
import subprocess

import pytest

from helpers import contar, html_de, numero, um

ABAS = ["os", "estoque", "financeiro", "fornecedores"]
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _openpyxl(dados):
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(dados), data_only=True)


def _todas_as_celulas(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if v is not None:
                    yield ws.title, v


def _texto_pdf(dados: bytes) -> str | None:
    try:
        import pypdf
        return "\n".join((p.extract_text() or "") for p in pypdf.PdfReader(io.BytesIO(dados)).pages)
    except ImportError:
        pass
    exe = shutil.which("pdftotext")
    if exe:
        r = subprocess.run([exe, "-", "-"], input=dados, capture_output=True, timeout=30)
        return r.stdout.decode("utf-8", "ignore")
    return None


def _checa_pdf(r, contexto):
    assert r.status_code == 200, f"{contexto}: HTTP {r.status_code}"
    assert r.mimetype == "application/pdf", f"{contexto}: Content-Type {r.mimetype}"
    d = r.get_data()
    assert d.startswith(b"%PDF"), f"{contexto}: nao comeca com %PDF"
    assert b"%%EOF" in d[-1024:], f"{contexto}: PDF truncado (sem %%EOF)"
    assert len(d) > 1500, f"{contexto}: PDF suspeitamente pequeno ({len(d)} bytes)"
    cd = r.headers.get("Content-Disposition", "")
    assert "filename" in cd.lower(), f"{contexto}: sem nome de arquivo no Content-Disposition ({cd!r})"
    return d


def _checa_xlsx(r, contexto):
    assert r.status_code == 200, f"{contexto}: HTTP {r.status_code}"
    assert r.mimetype == XLSX_MIME, f"{contexto}: Content-Type {r.mimetype}"
    assert r.get_data()[:2] == b"PK", f"{contexto}: nao e um zip/xlsx"
    assert ".xlsx" in r.headers.get("Content-Disposition", "").lower()
    return _openpyxl(r.get_data())


# --------------------------------------------------------------------------------------------- relatorios

@pytest.mark.parametrize("aba", ABAS)
def test_relatorio_pdf(admin, aba):
    r = admin.get(f"/relatorios/{aba}/exportar.pdf?periodo=12m")
    d = _checa_pdf(r, f"PDF {aba}")
    txt = _texto_pdf(d)
    if txt is not None:
        assert "TechMaint" in txt, "Cabecalho TechMaint no PDF (SPEC §4)"
        assert str(datetime.date.today().year) in txt, "Data de emissao no PDF (SPEC §4)"


@pytest.mark.parametrize("aba", ABAS)
def test_relatorio_xlsx(admin, aba):
    wb = _checa_xlsx(admin.get(f"/relatorios/{aba}/exportar.xlsx?periodo=12m"), f"XLSX {aba}")
    celulas = [str(v) for _, v in _todas_as_celulas(wb)]
    assert any("TechMaint" in c for c in celulas), "Cabecalho TechMaint na planilha (SPEC §4)"
    assert any(re.search(r"\d{2}/\d{2}/\d{4}", c) or isinstance(v, (datetime.date, datetime.datetime))
               for c, (_, v) in zip(celulas, _todas_as_celulas(wb))), "Data de emissao na planilha (SPEC §4)"
    assert any("12m" in c or "12 meses" in c.lower() or "ano" in c.lower() or "período" in c.lower() or "periodo" in c.lower() for c in celulas), \
        "Filtros aplicados devem constar na planilha (SPEC §4)"
    assert len(celulas) > 10, "Planilha praticamente vazia"


@pytest.mark.parametrize("aba", ABAS)
def test_relatorio_csv(admin, aba):
    r = admin.get(f"/relatorios/{aba}/exportar.csv?periodo=12m")
    assert r.status_code == 200
    assert r.mimetype in ("text/csv", "application/csv"), f"CSV com Content-Type {r.mimetype}"
    assert ".csv" in r.headers.get("Content-Disposition", "").lower()
    texto = r.get_data().decode("utf-8-sig")        # tem que decodificar como UTF-8 (com ou sem BOM)
    assert "�" not in texto, "Caracteres quebrados (encoding)"
    amostra = texto.splitlines()[:50]
    dialeto = csv.Sniffer().sniff("\n".join(amostra), delimiters=";,\t")
    linhas = list(csv.reader(io.StringIO(texto), dialeto))
    assert len(linhas) >= 2, "CSV sem cabecalho + dados"
    assert len({len(x) for x in linhas[1:] if x}) <= 3, "Linhas do CSV com numero muito irregular de colunas"


def test_csv_preserva_acentos(admin):
    r = admin.get("/relatorios/os/exportar.csv?periodo=12m")
    t = r.get_data().decode("utf-8-sig")
    assert "Técnico" in t or "Concluída" in t or "Maria Técnica" in t, "Acentos devem sobreviver ao CSV"


def test_csv_neutraliza_formula_injection(admin):
    """Celulas que comecam com = + - @ podem executar formulas no Excel (CSV injection)."""
    from helpers import executa
    executa("UPDATE pecas SET nome='=HYPERLINK(\"http://evil\",\"x\")' WHERE id=(SELECT MIN(id) FROM pecas)")
    for fmt in ("csv", "xlsx"):
        r = admin.get(f"/relatorios/estoque/exportar.{fmt}")
        assert r.status_code == 200
        if fmt == "csv":
            for linha in r.get_data().decode("utf-8-sig").splitlines():
                for cel in re.split(r"[;,\t]", linha):
                    assert not cel.strip('"').startswith("=HYPERLINK"), "CSV injection: celula iniciando com '='"
        else:
            wb = _openpyxl(r.get_data())
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for c in row:
                        assert not (c.data_type == "f" and "HYPERLINK" in str(c.value)), "XLSX injection: formula gravada"


def test_xlsx_da_aba_os_tem_as_mesmas_linhas_da_api(admin):
    j = admin.get("/api/relatorios/os?periodo=12m").get_json()
    wb = _checa_xlsx(admin.get("/relatorios/os/exportar.xlsx?periodo=12m"), "XLSX os")
    ws = wb.worksheets[0]
    cab = [str(c).lower() for c in j["tabela"]["colunas"]]
    achou_cab = None
    for row in ws.iter_rows(values_only=True):
        cels = [str(x).lower() for x in row if x is not None]
        if cels[:len(cab)] == cab:
            achou_cab = row
            break
    assert achou_cab, f"Cabecalho da tabela ({cab}) nao encontrado na planilha"
    texto = " ".join(str(v) for _, v in _todas_as_celulas(wb))
    for ln in j["tabela"]["linhas"][:5]:
        assert str(ln[0]) in texto and str(ln[1]) in texto, f"linha {ln} da API ausente da planilha"


def test_xlsx_valores_monetarios_sao_numericos(admin):
    wb = _checa_xlsx(admin.get("/relatorios/financeiro/exportar.xlsx?periodo=12m"), "XLSX financeiro")
    numericos = [v for _, v in _todas_as_celulas(wb) if isinstance(v, (int, float))]
    assert numericos, "Nenhum valor numerico na planilha financeira (valores como texto nao somam no Excel)"


def test_exportacao_respeita_filtros(admin):
    tudo = admin.get("/relatorios/os/exportar.csv?periodo=12m").get_data().decode("utf-8-sig").splitlines()
    did = um("SELECT id FROM departamentos WHERE nome='Logística'")
    filtrado = admin.get(f"/relatorios/os/exportar.csv?periodo=12m&departamento_id={did}").get_data().decode("utf-8-sig").splitlines()
    assert len(filtrado) <= len(tudo)
    sete = admin.get("/relatorios/os/exportar.csv?periodo=7d").get_data().decode("utf-8-sig").splitlines()
    assert len(sete) <= len(tudo)


def test_exportacao_csv_bate_com_a_tabela_da_tela(admin):
    j = admin.get("/api/relatorios/os?periodo=6m").get_json()
    t = admin.get("/relatorios/os/exportar.csv?periodo=6m").get_data().decode("utf-8-sig")
    linhas = [l for l in t.splitlines() if l.strip()]
    assert len(linhas) >= len(j["tabela"]["linhas"]) + 1, "CSV tem menos linhas que a tabela da tela"


@pytest.mark.parametrize("caminho", ["/relatorios/inexistente/exportar.pdf", "/relatorios/os/exportar.docx", "/relatorios/os/exportar.exe",
                                     "/relatorios/%27%3BDROP/exportar.csv", "/relatorios/..%2F..%2Fapp/exportar.csv"])
def test_formatos_e_abas_invalidos(admin, caminho):
    r = admin.get(caminho)
    assert r.status_code in (400, 404), f"{caminho} -> {r.status_code}"


@pytest.mark.parametrize("fmt", ["pdf", "xlsx", "csv"])
@pytest.mark.parametrize("qs", ["periodo=xx", "departamento_id=abc", "tecnico_id=99999", "departamento_id=99999&periodo=7d"])
def test_exportacao_com_filtros_invalidos_ou_sem_resultado(admin, fmt, qs):
    r = admin.get(f"/relatorios/os/exportar.{fmt}?{qs}")
    assert r.status_code in (200, 400), f"{fmt} {qs} -> {r.status_code}"
    if r.status_code == 200 and fmt == "pdf":
        assert r.get_data().startswith(b"%PDF")
    if r.status_code == 200 and fmt == "xlsx":
        _openpyxl(r.get_data())


def test_pdf_com_dados_hostis_e_unicode(admin):
    from helpers import executa
    executa("UPDATE pecas SET nome='<b>Peça</b> & \"aspas\" ç ã ñ 日本語 🙂' WHERE id=(SELECT MIN(id) FROM pecas)")
    r = admin.get("/relatorios/estoque/exportar.pdf")
    assert r.status_code == 200, "PDF quebrou com caracteres especiais (reportlab/markup)"
    assert r.get_data().startswith(b"%PDF")
    r = admin.get("/relatorios/estoque/exportar.xlsx")
    assert r.status_code == 200
    _openpyxl(r.get_data())


def test_pdf_com_muitas_linhas_pagina(admin):
    from helpers import executa
    for i in range(300):
        executa("INSERT INTO pecas (id_equipamento, nome, ncm, valor_unitario, qtde_estoque, estoque_minimo) "
                "VALUES ((SELECT MIN(id) FROM equipamentos), ?, '12345678', 1, 10, 5)", f"Peca em massa {i}")
    r = admin.get("/relatorios/estoque/exportar.pdf")
    d = _checa_pdf(r, "PDF com 300+ linhas")
    txt = _texto_pdf(d)
    if txt:
        assert "Peca em massa 299" in txt, "PDF nao paginou: linhas finais cortadas"


# --------------------------------------------------------------------------------------------- dashboard

def test_dashboard_pdf(admin):
    d = _checa_pdf(admin.get("/dashboard/exportar.pdf?periodo=12m"), "Dashboard PDF")
    txt = _texto_pdf(d)
    if txt:
        assert "TechMaint" in txt


def test_dashboard_xlsx(admin):
    wb = _checa_xlsx(admin.get("/dashboard/exportar.xlsx?periodo=12m"), "Dashboard XLSX")
    celulas = " ".join(str(v) for _, v in _todas_as_celulas(wb))
    for trecho in ("TechMaint",):
        assert trecho in celulas
    assert len(celulas) > 100


def test_dashboard_exportacao_respeita_filtros(admin):
    did = um("SELECT id FROM departamentos WHERE nome='Produção'")
    r = admin.get(f"/dashboard/exportar.xlsx?periodo=6m&departamento_id={did}")
    assert r.status_code == 200
    celulas = " ".join(str(v) for _, v in _todas_as_celulas(_openpyxl(r.get_data())))
    assert "Produção" in celulas, "Filtro de departamento aplicado deve constar na exportacao (SPEC §4)"


@pytest.mark.parametrize("fmt", ["pdf", "xlsx"])
@pytest.mark.parametrize("qs", ["periodo=xx", "departamento_id=abc", "equipamento_id=99999"])
def test_dashboard_exportacao_filtros_invalidos(admin, fmt, qs):
    r = admin.get(f"/dashboard/exportar.{fmt}?{qs}")
    assert r.status_code in (200, 400), f"{fmt} {qs} -> {r.status_code}"


# --------------------------------------------------------------------------------------------- permissao

@pytest.mark.parametrize("url", ["/relatorios/os/exportar.pdf", "/relatorios/os/exportar.xlsx", "/relatorios/os/exportar.csv",
                                 "/dashboard/exportar.pdf", "/dashboard/exportar.xlsx"])
def test_gestor_pode_exportar(gestor, url):
    assert gestor.get(url).status_code == 200


@pytest.mark.parametrize("url", ["/relatorios/os/exportar.pdf", "/relatorios/os/exportar.xlsx", "/relatorios/os/exportar.csv",
                                 "/dashboard/exportar.pdf", "/dashboard/exportar.xlsx", "/relatorios/financeiro/exportar.csv"])
def test_tecnico_e_comum_nao_exportam(tecnico, comum, url):
    assert tecnico.get(url).status_code == 403
    assert comum.get(url).status_code == 403


def test_exportacao_exige_login(anonimo):
    r = anonimo.get("/relatorios/os/exportar.pdf")
    assert r.status_code == 302
